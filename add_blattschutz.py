import openpyxl
from openpyxl.styles import Protection

def make_non_yellow_cells_readonly(file_path: str, output_path: str, password: str = None):
    # Excel-Datei laden
    wb = openpyxl.load_workbook(file_path)
    
    # Bekannte Gelb-Töne als ARGB/RGB-Hexcode (case-insensitive)
    YELLOW_CODES = {
        "FFFF00", 
    }
    
    # Schutz-Objekte für geschützte und freigegebene Zellen
    locked_style = Protection(locked=True)
    unlocked_style = Protection(locked=False)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Blattschutz aktivieren
        ws.protection.sheet = True
        if password:
            ws.protection.password = password

        # Alle verwendeten Zellen durchlaufen
        for row in ws.iter_rows():
            for cell in row:
                is_yellow = False
                
                # Prüfen, ob die Zelle eine Farbfüllung hat
                fill = cell.fill
                if fill and fill.fill_type == "solid" and fill.start_color:
                    color = fill.start_color
                    # Überprüfung auf Standard-RGB-Hex-Farbe
                    if color.type == "rgb" and color.rgb:
                        clean_rgb = color.rgb.upper()
                        # Manchmal liefert openpyxl 8 Stellen (ARGB) oder 6 Stellen (RGB)
                        if clean_rgb in YELLOW_CODES or clean_rgb[-6:] == "FFFF00":
                            is_yellow = True

                # Zelle entsprechend sperren oder entsperren
                if is_yellow:
                    cell.protection = unlocked_style
                else:
                    cell.protection = locked_style

    # Bearbeitete Datei speichern
    wb.save(output_path)
    print(f"Datei erfolgreich gespeichert unter: {output_path}")

# --- Beispielaufruf ---
if __name__ == "__main__":
    input_file = "test.xlsx"
    output_file = "test_geschuetzt.xlsx"
    
    # Passwort optional angeben (z. B. password="geheim")
    make_non_yellow_cells_readonly(input_file, output_file, password="abc")